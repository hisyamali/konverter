"""Logika baca file, pratinjau, normalisasi data, dan konversi PDF/CSV ke Excel."""

import csv
import math
import os
import re
import time
from concurrent.futures import ProcessPoolExecutor, TimeoutError, as_completed
from concurrent.futures.process import BrokenProcessPool
from decimal import Decimal, InvalidOperation

import pandas as pd
import pdfplumber
from flask import current_app, has_app_context
from openpyxl import Workbook

try:
    import psutil
except ImportError:
    psutil = None


def _get_memory_usage_mb():
    """Ambil penggunaan memori proses saat ini dalam MB."""
    if psutil is None:
        return None
    try:
        process = psutil.Process(os.getpid())
        mem = process.memory_info()
        return round(mem.rss / 1024 / 1024, 1)
    except Exception:
        return None


def _get_system_memory_percent():
    """Ambil persentase memori sistem yang terpakai."""
    if psutil is None:
        return None
    try:
        return psutil.virtual_memory().percent
    except Exception:
        return None


def _log_milestone(milestone, **context):
    """Log milestone konversi dengan context standar."""
    if has_app_context():
        current_app.logger.info(f"Conversion {milestone}", extra={"context": context})


def _log_warning(message, **context):
    """Log peringatan konversi (mis. degradasi mode) dengan context standar."""
    if has_app_context():
        current_app.logger.warning(message, extra={"context": context})


def _log_error(error, **context):
    """Log error konversi dengan context standar dan traceback."""
    if has_app_context():
        current_app.logger.error(
            "Conversion error",
            extra={"context": {"error": str(error), "error_type": type(error).__name__, **context}},
            exc_info=True,
        )


def _physical_core_count():
    """Hitung jumlah core fisik (Windows) agar worker tidak berlebih dari CPU nyata."""
    try:
        if os.name != "nt":
            return None

        import ctypes

        class SystemLogicalProcessorInformation(ctypes.Structure):
            _fields_ = [
                ("ProcessorMask", ctypes.c_size_t),
                ("Relationship", ctypes.c_int),
                ("Reserved", ctypes.c_ubyte * 16),
            ]

        get_info = ctypes.windll.kernel32.GetLogicalProcessorInformation
        length = ctypes.c_ulong(0)
        get_info(None, ctypes.byref(length))

        if length.value <= 0:
            return None

        count = length.value // ctypes.sizeof(SystemLogicalProcessorInformation)
        buffer = (SystemLogicalProcessorInformation * count)()

        if not get_info(buffer, ctypes.byref(length)):
            return None

        cores = sum(1 for item in buffer if item.Relationship == 0)
        return cores or None
    except Exception:
        return None


def _physical_cores():
    """Perkirakan core fisik; fallback ke separuh CPU logis bila deteksi gagal."""
    detected = _physical_core_count()

    if detected:
        return detected

    logical = os.cpu_count() or 1
    return max(1, logical // 2)


PARALLEL_MIN_PAGES = 12
PARALLEL_MAX_WORKERS_CAP = int(os.environ.get("KONVERTER_MAX_WORKERS_CAP", "8"))
PARALLEL_CHUNK_PAGES = int(os.environ.get("KONVERTER_CHUNK_PAGES", "20"))
PARALLEL_MAX_TASKS_PER_CHILD = int(os.environ.get("KONVERTER_TASKS_PER_CHILD", "32"))
PARALLEL_STALL_REPORT_SECONDS = 45
PARALLEL_DEGRADE_SECONDS = 90
PARALLEL_MEMORY_GB_PER_WORKER = float(os.environ.get("KONVERTER_MEM_GB_PER_WORKER", "0.15"))
PARALLEL_MEMORY_RESERVE_GB = float(os.environ.get("KONVERTER_MEM_RESERVE_GB", "0.4"))
PARALLEL_MEMORY_FLOOR_GB = float(os.environ.get("KONVERTER_MEM_FLOOR_GB", "0.35"))
POLITE_MODE = os.environ.get("KONVERTER_POLITE", "1") not in ("0", "false", "False")

_workers_override = int(os.environ.get("KONVERTER_MAX_WORKERS", "0") or "0")
PARALLEL_MAX_WORKERS = (
    _workers_override
    if _workers_override > 0
    else min(_physical_cores(), PARALLEL_MAX_WORKERS_CAP)
)

_BELOW_NORMAL_PRIORITY_CLASS = 0x00004000

PDF_TABLE_SETTINGS = {
    "vertical_strategy": "lines",
    "horizontal_strategy": "lines",
    "snap_tolerance": 5,
    "join_tolerance": 5,
    "edge_min_length": 5,
}

# Kolom teks untuk laporan mutasi bank berbasis teks (tanpa garis tabel), mis. Sumsel/BRI.
TEXT_DATE_X_MAX = 50
TEXT_DESC_X_MIN = 90
TEXT_DESC_X_MAX = 270
TEXT_DEBET_X = (270, 340)
TEXT_KREDIT_X = (340, 420)
TEXT_SALDO_X = (420, 510)
TEXT_REF_X = (510, 590)
TEXT_HEADER = ["Tanggal", "Keterangan", "Debet", "Kredit", "Saldo", "No. Referensi"]
TEXT_ROW_Y_TOLERANCE = 3

_WORKER_PDF = None


def convert_csv_to_excel(input_path, output_path, progress=None):
    """Baca CSV, normalisasi angka nominal, lalu tulis sebagai file Excel."""
    mem_start = _get_memory_usage_mb()
    _log_milestone("reading", input_path=input_path, memory_mb=mem_start)
    _report(progress, 12, "Membaca file CSV")
    dataframe = _read_csv(input_path)

    if dataframe.empty:
        raise ValueError("File CSV tidak memiliki data.")

    _log_milestone("normalize", rows=len(dataframe), columns=list(dataframe.columns))
    _report(progress, 70, "Menyusun kolom data")

    _normalize_nominal_values(dataframe)

    _log_milestone("write", rows=len(dataframe))
    _report(progress, 88, "Menulis file Excel")
    _write_excel(dataframe, output_path)

    mem_end = _get_memory_usage_mb()
    _log_milestone("done", output_path=output_path, memory_mb=mem_end, memory_delta_mb=round(mem_end - mem_start, 1) if mem_start and mem_end else None)
    _report(progress, 100, "Selesai")


def _write_excel(dataframe, output_path, sheet_name="Sheet1", include_header=True):
    """Tulis DataFrame ke workbook Excel write-only agar hemat memori untuk file besar."""
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(sheet_name)

    if include_header:
        sheet.append(list(dataframe.columns))

    for row in dataframe.itertuples(index=False, name=None):
        sheet.append(row)

    workbook.save(output_path)


def convert_pdf_to_excel(input_path, output_path, progress=None, total_pages=None):
    """Ekstrak semua tabel/transaksi PDF, rapikan header berulang, lalu ekspor ke Excel."""
    mem_start = _get_memory_usage_mb()
    sys_mem_start = _get_system_memory_percent()

    strategy = _detect_pdf_strategy(input_path)
    _log_milestone("reading", input_path=input_path, total_pages=total_pages, memory_mb=mem_start, system_memory_percent=sys_mem_start, strategy=strategy)

    if total_pages is None:
        _report(progress, 6, "Membuka file PDF")

    if strategy == "lines":
        builder = _RowBuilder()
        _collect_pdf_tables(input_path, progress, total_pages, builder.add_tables)
        rows = builder.rows
    else:
        rows = []
        _collect_text_rows(input_path, progress, total_pages, rows.extend)
        if rows:
            rows = [list(TEXT_HEADER)] + rows

    if not rows:
        raise ValueError("Tidak ada tabel yang bisa dibaca dari PDF.")

    mem_after_read = _get_memory_usage_mb()
    _log_milestone("normalize", rows=len(rows), columns=None, memory_mb=mem_after_read, memory_delta_mb=round(mem_after_read - mem_start, 1) if mem_start and mem_after_read else None, strategy=strategy)
    _report(progress, 86, "Menyusun kolom data")

    _normalize_table(rows)
    dataframe = pd.DataFrame(rows)
    _normalize_nominal_values(dataframe)

    _log_milestone("write", rows=len(dataframe), columns=list(dataframe.columns))
    _report(progress, 93, "Menulis file Excel")
    _write_excel(dataframe, output_path, sheet_name="Data", include_header=False)

    mem_end = _get_memory_usage_mb()
    sys_mem_end = _get_system_memory_percent()
    _log_milestone("done", output_path=output_path, rows=len(dataframe), memory_mb=mem_end, memory_delta_mb=round(mem_end - mem_start, 1) if mem_start and mem_end else None, system_memory_percent=sys_mem_end, strategy=strategy)
    _report(progress, 100, "Selesai")


def _collect_pdf_tables(input_path, progress, total_pages, on_tables):
    """Baca tabel PDF sekuensial atau paralel, kirim ke on_tables sesuai urutan halaman."""
    if total_pages is None:
        total_pages = _count_pages(input_path)

    workers = _worker_count(total_pages)

    _report(progress, 8, f"Halaman 0 dari halaman {total_pages}")

    if workers == 1:
        _read_pages_sequential(input_path, total_pages, progress, on_tables)
        return

    _read_chunks_parallel(input_path, _build_chunks(total_pages), workers, total_pages, progress, on_tables)


def _count_pages(input_path):
    """Hitung jumlah halaman PDF sebelum proses ekstraksi berjalan."""
    with pdfplumber.open(input_path) as pdf:
        return len(pdf.pages)


def _worker_count(total_pages):
    """Tentukan jumlah worker paralel berdasarkan halaman, CPU, dan memori kosong."""
    if total_pages <= PARALLEL_MIN_PAGES:
        return 1

    logical = os.cpu_count() or 1
    workers = min(logical, PARALLEL_MAX_WORKERS)
    available_gb = _available_memory_gb()

    if available_gb is not None:
        by_memory = int((available_gb - PARALLEL_MEMORY_RESERVE_GB) / PARALLEL_MEMORY_GB_PER_WORKER)
        workers = min(workers, max(by_memory, 1))

    return max(1, workers)


def _memory_pressure():
    """Cek apakah RAM kosong turun di bawah ambang aman saat konversi berjalan."""
    available_gb = _available_memory_gb()
    return available_gb is not None and available_gb < PARALLEL_MEMORY_FLOOR_GB


def _available_memory_gb():
    """Ambil estimasi RAM tersedia di Windows atau Linux untuk membatasi worker."""
    try:
        if os.name == "nt":
            import ctypes

            class MemoryStatusEx(ctypes.Structure):
                _fields_ = [
                    ("length", ctypes.c_ulong),
                    ("memory_load", ctypes.c_ulong),
                    ("total_phys", ctypes.c_ulonglong),
                    ("avail_phys", ctypes.c_ulonglong),
                    ("total_page_file", ctypes.c_ulonglong),
                    ("avail_page_file", ctypes.c_ulonglong),
                    ("total_virtual", ctypes.c_ulonglong),
                    ("avail_virtual", ctypes.c_ulonglong),
                    ("avail_ext_virtual", ctypes.c_ulonglong),
                ]

            status = MemoryStatusEx()
            status.length = ctypes.sizeof(MemoryStatusEx)
            ctypes.windll.kernel32.GlobalMemoryStatusEx(ctypes.byref(status))
            return status.avail_phys / (1024 ** 3)

        if os.path.exists("/proc/meminfo"):
            with open("/proc/meminfo", "r", encoding="utf-8") as meminfo:
                for line in meminfo:
                    if line.startswith("MemAvailable:"):
                        return int(line.split()[1]) / (1024 ** 2)

        return None
    except Exception:
        return None


def _build_chunks(total_pages):
    """Bagi halaman PDF menjadi rentang kecil untuk diproses worker paralel."""
    return [
        (start, min(start + PARALLEL_CHUNK_PAGES, total_pages))
        for start in range(0, total_pages, PARALLEL_CHUNK_PAGES)
    ]


def _read_pages_sequential(input_path, total_pages, progress, on_tables, start=0):
    """Baca halaman PDF satu per satu, digunakan untuk file kecil atau fallback aman."""
    pages_read = start

    with pdfplumber.open(input_path) as pdf:
        for page in pdf.pages[start:]:
            on_tables(_page_tables(page))
            pages_read += 1
            _report(progress, _percent_from_pages(pages_read, total_pages), _status_message(pages_read, total_pages))


def _read_chunks_parallel(input_path, chunks, workers, total_pages, progress, on_tables):
    """Baca chunk PDF dengan banyak proses sambil mengalirkan hasil sesuai urutan halaman."""
    chunk_end = dict(chunks)
    buffered = {}
    next_expected = 0
    pages_read = 0
    last_success = time.monotonic()
    degraded = False
    degrade_reason = None

    executor = ProcessPoolExecutor(
        max_workers=workers,
        initializer=_init_pdf_worker,
        initargs=(str(input_path),),
        max_tasks_per_child=PARALLEL_MAX_TASKS_PER_CHILD,
    )
    futures = {executor.submit(_extract_chunk_task, chunk): chunk for chunk in chunks}
    pending = set(futures)

    try:
        while pending and not degraded:
            try:
                for future in as_completed(pending, timeout=PARALLEL_STALL_REPORT_SECONDS):
                    chunk_start, chunk_tables = future.result()
                    buffered[chunk_start] = chunk_tables
                    chunk_start, chunk_stop = futures[future]
                    pages_read += chunk_stop - chunk_start
                    pending.remove(future)
                    last_success = time.monotonic()

                    while next_expected in buffered:
                        on_tables(buffered.pop(next_expected))
                        next_expected = chunk_end[next_expected]

                    _report(progress, _percent_from_pages(pages_read, total_pages), _status_message(pages_read, total_pages))

                    if _memory_pressure():
                        degraded = True
                        degrade_reason = "memori menipis"
                        break
            except TimeoutError:
                stalled_for = int(time.monotonic() - last_success)
                _report(
                    progress,
                    _percent_from_pages(pages_read, total_pages),
                    f"{_status_message(pages_read, total_pages)} — pembacaan melambat (±{stalled_for} detik), tetap berjalan",
                )

                if stalled_for >= PARALLEL_DEGRADE_SECONDS:
                    degraded = True
                    degrade_reason = "pembacaan melambat"
    except BrokenProcessPool as error:
        _abort_executor(executor, pending)
        _log_error(error, stage="parallel_read", total_pages=total_pages, pages_read=pages_read)
        raise RuntimeError("Proses pembacaan PDF terhenti. Silakan coba lagi.")
    except Exception as error:
        _abort_executor(executor, pending)
        _log_error(error, stage="parallel_read", total_pages=total_pages, pages_read=pages_read)
        raise
    else:
        if degraded:
            _abort_executor(executor, pending)
        else:
            executor.shutdown(wait=True)

    if degraded:
        if degrade_reason == "memori menipis":
            message = f"{_status_message(next_expected, total_pages)} — memori menipis, beralih ke mode hemat (satu-per-satu)"
        else:
            message = f"{_status_message(next_expected, total_pages)} — beralih ke mode satu-per-satu agar pasti selesai"

        _log_warning("Conversion degraded", reason=degrade_reason, next_expected=next_expected, total_pages=total_pages)
        _report(progress, _percent_from_pages(next_expected, total_pages), message)
        _read_pages_sequential(input_path, total_pages, progress, on_tables, start=next_expected)


def _abort_executor(executor, pending):
    """Batalkan pekerjaan paralel yang tersisa saat perlu fallback atau terjadi error."""
    for future in pending:
        future.cancel()

    try:
        executor.shutdown(wait=False, cancel_futures=True)
    except ValueError:
        pass


def _init_pdf_worker(pdf_path):
    """Buka PDF sekali di tiap worker agar setiap task tidak membuka file ulang."""
    _apply_polite_priority()
    global _WORKER_PDF
    _WORKER_PDF = pdfplumber.open(pdf_path)


def _apply_polite_priority():
    """Turunkan prioritas proses worker agar aplikasi lain tetap responsif (Windows)."""
    if not POLITE_MODE:
        return

    try:
        if os.name == "nt":
            import ctypes

            handle = ctypes.windll.kernel32.GetCurrentProcess()
            ctypes.windll.kernel32.SetPriorityClass(handle, _BELOW_NORMAL_PRIORITY_CLASS)
        else:
            os.nice(5)
    except Exception:
        pass


def _extract_chunk_task(chunk):
    """Wrapper task worker untuk mengembalikan posisi chunk dan hasil tabelnya."""
    start, end = chunk
    return start, _extract_chunk(_WORKER_PDF, start, end)


def _detect_pdf_strategy(input_path):
    """Deteksi strategi ekstraksi: 'lines' jika halaman pertama punya tabel bergaris dengan data, selain itu 'text'."""
    try:
        with pdfplumber.open(input_path) as pdf:
            if not pdf.pages:
                return "lines"
            preview = pdf.pages[0].extract_tables(PDF_TABLE_SETTINGS)
    except Exception:
        return "lines"

    data_rows = 0
    if preview:
        for table in preview:
            if table:
                data_rows += max(len(table) - 1, 0)

    return "lines" if data_rows >= 1 else "text"


def _extract_text_page(page):
    """Ekstrak transaksi dari satu halaman PDF berbasis teks (posisi kolom tetap)."""
    words = page.extract_words()
    words.sort(key=lambda w: (w["top"], w["x0"]))

    clusters = []
    for word in words:
        if clusters and abs(clusters[-1]["y"] - word["top"]) <= TEXT_ROW_Y_TOLERANCE:
            clusters[-1]["words"].append(word)
        else:
            clusters.append({"y": word["top"], "words": [word]})

    date_re = re.compile(r"^\d{2}-\d{2}-\d{4}$")
    rows = []
    current = None

    def _amount_in(tokens, bounds):
        low, high = bounds
        return " ".join(w["text"] for w in tokens if low <= w["x0"] < high)

    for cluster in clusters:
        tokens = sorted(cluster["words"], key=lambda w: w["x0"])
        if not tokens:
            continue

        first = tokens[0]
        if date_re.match(first["text"]) and first["x0"] < TEXT_DATE_X_MAX:
            if current is not None:
                rows.append(_finalize_text_transaction(current))
            current = {
                "date": first["text"],
                "debet": _amount_in(tokens, TEXT_DEBET_X),
                "kredit": _amount_in(tokens, TEXT_KREDIT_X),
                "saldo": _amount_in(tokens, TEXT_SALDO_X),
                "ref": _amount_in(tokens, TEXT_REF_X),
                "desc": [],
            }
            continue

        if current is None:
            continue

        line_text = " ".join(w["text"] for w in tokens)
        if re.match(r"Page:", line_text):
            continue

        desc_words = [w["text"] for w in tokens if TEXT_DESC_X_MIN <= w["x0"] < TEXT_DESC_X_MAX]
        if desc_words:
            current["desc"].append(" ".join(desc_words))

    if current is not None:
        rows.append(_finalize_text_transaction(current))

    return rows


def _finalize_text_transaction(transaction):
    description = "\n".join(transaction["desc"]) if transaction["desc"] else ""
    return [
        transaction["date"],
        description,
        transaction["debet"],
        transaction["kredit"],
        transaction["saldo"],
        transaction["ref"],
    ]


def _extract_text_chunk_task(chunk):
    """Task worker: ekstrak transaksi teks dari rentang halaman PDF."""
    start, end = chunk
    rows = []
    for page in _WORKER_PDF.pages[start:end]:
        rows.extend(_extract_text_page(page))
    return start, rows


def _read_text_pages_sequential(input_path, total_pages, progress, on_rows, start=0):
    """Baca halaman PDF berbasis teks satu per satu."""
    pages_read = start
    with pdfplumber.open(input_path) as pdf:
        for page in pdf.pages[start:]:
            on_rows(_extract_text_page(page))
            pages_read += 1
            _report(progress, _percent_from_pages(pages_read, total_pages), _status_message(pages_read, total_pages))


def _read_text_chunks_parallel(input_path, chunks, workers, total_pages, progress, on_rows):
    """Baca chunk PDF berbasis teks dengan banyak proses, kirim hasil sesuai urutan halaman."""
    chunk_end = dict(chunks)
    buffered = {}
    next_expected = 0
    pages_read = 0
    last_success = time.monotonic()
    degraded = False
    degrade_reason = None

    executor = ProcessPoolExecutor(
        max_workers=workers,
        initializer=_init_pdf_worker,
        initargs=(str(input_path),),
        max_tasks_per_child=PARALLEL_MAX_TASKS_PER_CHILD,
    )
    futures = {executor.submit(_extract_text_chunk_task, chunk): chunk for chunk in chunks}
    pending = set(futures)

    try:
        while pending and not degraded:
            try:
                for future in as_completed(pending, timeout=PARALLEL_STALL_REPORT_SECONDS):
                    chunk_start, chunk_rows = future.result()
                    buffered[chunk_start] = chunk_rows
                    chunk_start_key, chunk_stop = futures[future]
                    pages_read += chunk_stop - chunk_start_key
                    pending.remove(future)
                    last_success = time.monotonic()

                    while next_expected in buffered:
                        on_rows(buffered.pop(next_expected))
                        next_expected = chunk_end[next_expected]

                    _report(progress, _percent_from_pages(pages_read, total_pages), _status_message(pages_read, total_pages))

                    if _memory_pressure():
                        degraded = True
                        degrade_reason = "memori menipis"
                        break
            except TimeoutError:
                stalled_for = int(time.monotonic() - last_success)
                _report(
                    progress,
                    _percent_from_pages(pages_read, total_pages),
                    f"{_status_message(pages_read, total_pages)} — pembacaan melambat (±{stalled_for} detik), tetap berjalan",
                )

                if stalled_for >= PARALLEL_DEGRADE_SECONDS:
                    degraded = True
                    degrade_reason = "pembacaan melambat"
    except BrokenProcessPool:
        _abort_executor(executor, pending)
        raise RuntimeError("Proses pembacaan PDF terhenti. Silakan coba lagi.")
    except Exception as error:
        _abort_executor(executor, pending)
        _log_error(error, stage="parallel_text_read", total_pages=total_pages, pages_read=pages_read)
        raise
    else:
        if degraded:
            _abort_executor(executor, pending)
        else:
            executor.shutdown(wait=True)

    if degraded:
        if degrade_reason == "memori menipis":
            message = f"{_status_message(next_expected, total_pages)} — memori menipis, beralih ke mode hemat (satu-per-satu)"
        else:
            message = f"{_status_message(next_expected, total_pages)} — beralih ke mode satu-per-satu agar pasti selesai"

        _log_warning("Conversion degraded", reason=degrade_reason, next_expected=next_expected, total_pages=total_pages)
        _report(progress, _percent_from_pages(next_expected, total_pages), message)
        _read_text_pages_sequential(input_path, total_pages, progress, on_rows, start=next_expected)


def _collect_text_rows(input_path, progress, total_pages, on_rows):
    """Baca transaksi teks PDF sekuensial atau paralel, kirim ke on_rows sesuai urutan halaman."""
    if total_pages is None:
        total_pages = _count_pages(input_path)

    workers = _worker_count(total_pages)

    _report(progress, 8, f"Halaman 0 dari halaman {total_pages}")

    if workers == 1:
        _read_text_pages_sequential(input_path, total_pages, progress, on_rows)
        return

    _read_text_chunks_parallel(input_path, _build_chunks(total_pages), workers, total_pages, progress, on_rows)


def _page_tables(page):
    """Ekstrak dan bersihkan semua tabel yang ditemukan pada satu halaman PDF."""
    tables = []

    try:
        for table in page.extract_tables(PDF_TABLE_SETTINGS):
            cleaned_table = _clean_table(table)

            if cleaned_table:
                tables.append(cleaned_table)
    finally:
        page.flush_cache()

    return tables


def _extract_chunk(pdf, start, end):
    """Ekstrak tabel dari rentang halaman PDF tertentu."""
    tables = []

    for page in pdf.pages[start:end]:
        tables.extend(_page_tables(page))

    return tables


def _percent_from_pages(pages_read, total_pages):
    """Ubah jumlah halaman terbaca menjadi persentase progress tahap ekstraksi."""
    return 8 + int(76 * pages_read / max(total_pages, 1))


def _status_message(pages_read, total_pages):
    """Buat pesan progress pembacaan halaman untuk frontend."""
    return f"Halaman {pages_read} dari halaman {total_pages}"


_PAGE_STATUS_RE = re.compile(r"^Halaman \d+ dari halaman \d+$")
_progress_state = {}


def _report(progress, percent, message):
    """Panggil callback progress jika pemanggil menyediakannya.

    Update "Halaman X dari Y" di-throttle agar hanya terkirim saat persentase
    naik minimal 1%, sementara pesan fase/error tetap selalu terkirim.
    """
    if progress is None:
        return

    key = id(progress)
    last_percent = _progress_state.get(key)

    is_page_status = bool(message and _PAGE_STATUS_RE.match(message))
    if is_page_status and last_percent is not None and int(percent) == last_percent:
        return

    _progress_state[key] = int(percent)
    progress(percent, message)


PREVIEW_SAMPLE_ROWS = 10
PREVIEW_PDF_TABLE_ROWS = 5


def preview_csv_file(path):
    """Buat ringkasan CSV: jumlah baris/kolom, header, dan sampel baris."""
    dataframe = _read_csv(path)

    if dataframe.empty:
        raise ValueError("File CSV tidak memiliki data.")

    headers = [str(column) for column in dataframe.columns]
    sample = dataframe.head(PREVIEW_SAMPLE_ROWS)

    return {
        "row_count": len(dataframe),
        "column_count": len(headers),
        "headers": headers,
        "sample_rows": [_json_rows(row) for row in sample.itertuples(index=False, name=None)],
    }


PREVIEW_TIMED_PAGES = 5


def preview_pdf_file(path):
    """Buat ringkasan PDF: jumlah halaman, sampel baris tabel, dan estimasi durasi."""
    with pdfplumber.open(path) as pdf:
        page_count = len(pdf.pages)

        if page_count == 0:
            raise ValueError("Tidak ada halaman yang bisa dibaca dari PDF.")

        strategy = _detect_pdf_strategy(path)

        sample_rows = []
        if strategy == "lines":
            preview = pdf.pages[0].extract_tables(PDF_TABLE_SETTINGS)
            if preview:
                for row in preview[0][:PREVIEW_PDF_TABLE_ROWS]:
                    sample_rows.append(
                        [(cell or "").strip() if cell else "" for cell in row]
                    )
        else:
            for row in _extract_text_page(pdf.pages[0])[:PREVIEW_PDF_TABLE_ROWS]:
                sample_rows.append([str(cell) for cell in row])

        if page_count > 1:
            page_count_timed = min(PREVIEW_TIMED_PAGES, page_count)
            start = min(page_count // 2, page_count - page_count_timed)
            timed_start = time.perf_counter()
            timed_pages = 0

            for index in range(start, start + page_count_timed):
                try:
                    if strategy == "lines":
                        pdf.pages[index].extract_tables(PDF_TABLE_SETTINGS)
                    else:
                        pdf.pages[index].extract_words()
                except Exception:
                    continue

                timed_pages += 1

            elapsed_sample = time.perf_counter() - timed_start
            per_page_seconds = elapsed_sample / max(timed_pages, 1)
        else:
            timed_pages = 0
            per_page_seconds = 0.0

    return {
        "page_count": page_count,
        "sample_rows": sample_rows,
        "estimated_seconds": int(round(per_page_seconds * page_count)) if timed_pages else 0,
    }


def _json_rows(row):
    """Konversi satu baris data menjadi list nilai yang aman untuk JSON."""
    return [_json_value(value) for value in row]


def _json_value(value):
    """Normalisasi nilai pandas/numpy agar bisa dikirim lewat JSON."""
    if value is None:
        return None

    if hasattr(value, "item"):
        value = value.item()

    if isinstance(value, bool):
        return value

    if isinstance(value, int):
        return value

    if isinstance(value, float):
        return value if math.isfinite(value) else None

    return str(value)


def _read_csv(input_path):
    """Baca CSV dengan parser khusus transaksi terlebih dahulu, lalu parser umum."""
    table = _read_transaction_csv(input_path)
    if table is not None:
        return table

    return _read_generic_csv(input_path)


def _read_generic_csv(input_path):
    """Baca CSV umum dengan deteksi delimiter dan fallback encoding latin-1."""
    if not _has_content(input_path):
        return pd.DataFrame()

    delimiter = _sniff_delimiter(input_path)

    if delimiter is not None:
        try:
            return pd.read_csv(input_path, sep=delimiter, engine="c")
        except Exception:
            try:
                return pd.read_csv(input_path, sep=delimiter, engine="c", encoding="latin-1")
            except Exception:
                pass

    try:
        return pd.read_csv(input_path, sep=None, engine="python")
    except UnicodeDecodeError:
        return pd.read_csv(input_path, sep=None, engine="python", encoding="latin-1")


def _has_content(input_path):
    """Cek apakah file berisi karakter non-whitespace sebelum dibaca pandas."""
    whitespace = b"\x00\r\n\t "
    try:
        with open(input_path, "rb") as file:
            while True:
                chunk = file.read(65536)

                if not chunk:
                    return False

                if any(byte not in whitespace for byte in chunk):
                    return True
    except OSError:
        return True


def _sniff_delimiter(input_path):
    """Deteksi delimiter koma atau titik koma dari sampel awal file."""
    try:
        with open(input_path, "rb") as file:
            sample = file.read(8192)
    except OSError:
        return None

    text = None

    for encoding in ("utf-8-sig", "latin-1"):
        try:
            text = sample.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if text is None:
        text = sample.decode("latin-1", errors="replace")

    comma_count = 0
    semicolon_count = 0

    for line in text.splitlines():
        stripped = line.strip()

        if not stripped:
            continue

        comma_count += _count_unquoted(stripped, ",")
        semicolon_count += _count_unquoted(stripped, ";")

    if comma_count == 0 and semicolon_count == 0:
        return None

    if comma_count == semicolon_count:
        return None

    return "," if comma_count > semicolon_count else ";"


def _count_unquoted(line, character):
    """Hitung separator di luar tanda kutip agar CSV quoted tidak salah deteksi."""
    count = 0
    in_quotes = False
    index = 0
    length = len(line)

    while index < length:
        current = line[index]

        if current == '"':
            if in_quotes and index + 1 < length and line[index + 1] == '"':
                index += 2
                continue

            in_quotes = not in_quotes
        elif current == character and not in_quotes:
            count += 1

        index += 1

    return count


def _read_transaction_csv(input_path):
    """Baca CSV laporan transaksi yang punya metadata sebelum header tabel utama."""
    lines = _read_text_lines(input_path)

    if not lines:
        return None

    header_index, delimiter, columns = _find_transaction_header(lines)

    if header_index is None:
        return None

    rows = []
    column_count = len(columns)

    for line in lines[header_index + 1:]:
        if not line.strip():
            continue

        row = next(csv.reader([line], delimiter=delimiter))
        row = _fit_csv_row(row, column_count)

        if any(cell.strip() for cell in row):
            rows.append(row)

    return pd.DataFrame(rows, columns=columns)


def _read_text_lines(input_path):
    """Baca seluruh file teks dengan encoding UTF-8 BOM atau latin-1."""
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            with open(input_path, encoding=encoding, newline="") as file:
                return file.read().splitlines()
        except UnicodeDecodeError:
            continue

    return []


def _find_transaction_header(lines):
    """Cari baris header transaksi berdasarkan kolom kunci seperti tanggal/keterangan."""
    for index, line in enumerate(lines):
        for delimiter in (",", ";"):
            columns = next(csv.reader([line], delimiter=delimiter))
            normalized_columns = _normalize_row(columns)

            if "tanggal" in normalized_columns and "keterangan" in normalized_columns:
                return index, delimiter, columns

    return None, None, None


def _fit_csv_row(row, column_count):
    """Sesuaikan jumlah kolom CSV, termasuk menggabungkan deskripsi yang terpecah."""
    row = [cell.strip() for cell in row]

    if len(row) == column_count:
        return row

    if len(row) > column_count and column_count >= 3:
        fixed_tail_count = column_count - 2
        description = " ".join(cell for cell in row[1:-fixed_tail_count] if cell).strip()
        return [row[0], description, *row[-fixed_tail_count:]]

    return row + [""] * (column_count - len(row))


def _clean_table(table):
    """Buang baris kosong dan rapikan spasi sel hasil ekstraksi PDF."""
    cleaned_rows = []

    for row in table:
        if row is None:
            continue

        cleaned_row = [(cell or "").strip() for cell in row]
        if any(cleaned_row):
            cleaned_rows.append(cleaned_row)

    return cleaned_rows


class _RowBuilder:
    """Gabungkan tabel PDF secara streaming tanpa menyimpan seluruh tabel di memori."""

    def __init__(self):
        self.header = None
        self.rows = []

    def add_tables(self, tables):
        for table in tables:
            if not table:
                continue

            if self.header is None:
                self.header = table[0]
                self.rows.extend(table)
            elif _same_row(table[0], self.header):
                self.rows.extend(table[1:])
            else:
                self.rows.extend(table)


def _append_table_without_duplicate_header(table, header, rows):
    """Gabungkan tabel PDF (in-place) sambil menghindari header berulang tiap halaman."""
    current_header = table[0]

    if header is None:
        rows.extend(table)
        return current_header, rows

    if _same_row(current_header, header):
        rows.extend(table[1:])
    else:
        rows.extend(table)

    return header, rows


def _same_row(first_row, second_row):
    """Bandingkan dua baris setelah case dan spasi dinormalisasi."""
    return _normalize_row(first_row) == _normalize_row(second_row)


def _normalize_row(row):
    """Ubah baris menjadi bentuk pembanding yang lowercase dan satu spasi."""
    return [" ".join(cell.lower().split()) for cell in row]


def _normalize_table(table):
    """Samakan panjang semua baris tabel agar bisa dibuat DataFrame (in-place)."""
    max_columns = max(len(row) for row in table)

    for row in table:
        missing = max_columns - len(row)

        if missing:
            row.extend([""] * missing)

    return table


_IDENTIFIER_PLAIN = re.compile(r"0\d+")
_IDENTIFIER_GROUPED = re.compile(r"0\d{1,2}(?:,\d{3})+")
_SIMPLE_INTEGER = re.compile(r"[+-]?\d+")
_TOKEN_SEPARATOR = re.compile(r"[\n\r\t ]")


def _normalize_nominal_values(dataframe):
    """Konversi kolom nominal ke angka tanpa mengubah kolom yang terlihat seperti ID."""
    for column in dataframe.columns:
        series = dataframe[column]

        if _has_identifier_value(series):
            continue

        dataframe[column] = _normalize_nominal_column(series)

    return dataframe


def _has_identifier_value(series):
    """Deteksi kolom berisi nomor identitas/rekening agar tidak diperlakukan angka nominal."""
    cleaned = _clean_series_token(series)
    return bool(
        cleaned.str.fullmatch(_IDENTIFIER_PLAIN.pattern, na=False).any()
        or cleaned.str.fullmatch(_IDENTIFIER_GROUPED.pattern, na=False).any()
    )


def _clean_series_token(series):
    """Bersihkan token kolom dari whitespace dan tanda debit/kredit untuk analisis angka."""
    cleaned = series.map(lambda value: "" if value is None else str(value))
    cleaned = cleaned.str.strip().str.replace(_TOKEN_SEPARATOR.pattern, "", regex=True).str.rstrip("CDcd")
    return cleaned.mask(cleaned.isin(["-", "--", ".", ","]), "")


def _normalize_nominal_column(series):
    """Normalisasi satu kolom nilai nominal sambil mempertahankan nilai non-angka."""
    values = series.tolist()
    cleaned = _clean_series_token(series).tolist()
    result = []

    for index, value in enumerate(values):
        token = cleaned[index]

        if _SIMPLE_INTEGER.fullmatch(token):
            number = _to_simple_integer(token)
            result.append(value if number is None else number)
        else:
            result.append(_normalize_cell(value))

    return result


def _to_simple_integer(token):
    """Konversi token integer sederhana, kembalikan None jika tidak valid."""
    try:
        return int(token)
    except (ValueError, TypeError):
        return None


def _normalize_cell(value):
    """Konversi satu sel nominal ke int/float jika formatnya dapat dikenali."""
    if value is None:
        return None

    token = _clean_number_token(value)

    if not token:
        return value

    number = _to_number(token)
    return value if number is None else number


def _clean_number_token(value):
    """Rapikan teks angka dari spasi, newline, dan suffix CD sebelum parsing."""
    token = str(value).strip()
    token = token.replace("\n", " ").replace("\r", " ").replace("\t", " ")
    token = token.replace(" ", "")
    token = token.rstrip("CDcd")

    if token in {"-", "--", ".", ","}:
        return ""

    return token


def _to_number(value):
    """Parse format angka Indonesia/umum menjadi int atau float."""
    token = _clean_number_token(value)

    if not token:
        return None

    lowered = token.lower()

    if "e" in lowered:
        mantissa, _, exponent = lowered.partition("e")

        try:
            return _to_integral(Decimal(mantissa) * (Decimal(10) ** int(exponent)))
        except (InvalidOperation, ValueError):
            return None

    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", token):
        token = token.replace(".", "").replace(",", ".")
    else:
        token = token.replace(",", "")

    try:
        return _to_integral(Decimal(token))
    except InvalidOperation:
        return None


def _to_integral(number):
    """Kembalikan int jika Decimal bulat, selain itu float."""
    integral = number.to_integral_value()
    return int(integral) if number == integral else float(number)
